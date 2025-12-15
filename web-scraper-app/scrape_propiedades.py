#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Scraper de Propiedades MercadoLibre - Arriendos Las Condes
==========================================================
Script para extraer listados de arriendos en Las Condes y exportar a Excel.
"""

import re
import time
import random
import logging
from datetime import datetime
from dataclasses import dataclass, asdict
from typing import List, Optional

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Configurar logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


@dataclass
class Propiedad:
    """Estructura de datos para una propiedad."""
    titulo: str
    precio: Optional[float]
    moneda: str
    ubicacion: str
    metros_cuadrados: Optional[int]
    dormitorios: Optional[int]
    banos: Optional[int]
    url: str
    imagen_url: Optional[str]
    tipo_propiedad: Optional[str]
    gastos_comunes: Optional[float]
    fecha_scraping: str


class MercadoLibrePropiedadesScraper:
    """
    Scraper para propiedades de MercadoLibre Chile.
    Específico para arriendos en Las Condes.
    """

    BASE_URL = "https://inmuebles.mercadolibre.cl"

    # User agents para rotación
    USER_AGENTS = [
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:121.0) Gecko/20100101 Firefox/121.0',
        'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.1 Safari/605.1.15',
    ]

    def __init__(self, delay_min: float = 2, delay_max: float = 4, max_retries: int = 3):
        self.delay_min = delay_min
        self.delay_max = delay_max
        self.max_retries = max_retries
        self.session = requests.Session()
        self.propiedades: List[Propiedad] = []

    def get_headers(self) -> dict:
        """Genera headers con User-Agent aleatorio."""
        return {
            'User-Agent': random.choice(self.USER_AGENTS),
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'es-CL,es;q=0.9,en;q=0.8',
            'Accept-Encoding': 'gzip, deflate, br',
            'Connection': 'keep-alive',
            'Upgrade-Insecure-Requests': '1',
        }

    def random_delay(self):
        """Aplica un delay aleatorio entre requests."""
        delay = random.uniform(self.delay_min, self.delay_max)
        logger.debug(f"Esperando {delay:.2f} segundos...")
        time.sleep(delay)

    def fetch_page(self, url: str) -> Optional[BeautifulSoup]:
        """Obtiene y parsea una página con retry logic."""
        for attempt in range(1, self.max_retries + 1):
            try:
                logger.info(f"Obteniendo URL (intento {attempt}): {url}")
                response = self.session.get(url, headers=self.get_headers(), timeout=30)
                response.raise_for_status()
                return BeautifulSoup(response.content, 'lxml')

            except requests.exceptions.RequestException as e:
                logger.warning(f"Error en intento {attempt}: {e}")
                if attempt < self.max_retries:
                    time.sleep(2 ** attempt)

        logger.error(f"Falló después de {self.max_retries} intentos")
        return None

    def build_search_url(self, page: int = 1) -> str:
        """
        Construye la URL de búsqueda para arriendos en Las Condes.
        """
        # URL para arriendos en Las Condes, Región Metropolitana
        base = f"{self.BASE_URL}/arriendo/las-condes-metropolitana"

        if page == 1:
            return base
        else:
            # MercadoLibre usa _Desde_ para paginación (48 items por página)
            offset = (page - 1) * 48 + 1
            return f"{base}_Desde_{offset}"

    def clean_price(self, price_text: str) -> tuple:
        """
        Extrae precio y moneda del texto.
        Retorna (precio, moneda).
        """
        if not price_text:
            return None, "CLP"

        price_text = price_text.strip()

        # Detectar moneda
        if 'UF' in price_text.upper():
            moneda = "UF"
        elif 'US' in price_text.upper() or 'USD' in price_text.upper():
            moneda = "USD"
        else:
            moneda = "CLP"

        # Extraer números
        numbers = re.findall(r'[\d.,]+', price_text)
        if numbers:
            price_str = numbers[0].replace('.', '').replace(',', '.')
            try:
                return float(price_str), moneda
            except ValueError:
                pass

        return None, moneda

    def extract_number(self, text: str) -> Optional[int]:
        """Extrae el primer número entero de un texto."""
        if not text:
            return None
        numbers = re.findall(r'\d+', text)
        if numbers:
            return int(numbers[0])
        return None

    def parse_property(self, card) -> Optional[Propiedad]:
        """Parsea una tarjeta de propiedad."""
        try:
            # Título
            title_elem = (
                card.select_one('h2.poly-box.poly-component__title a') or
                card.select_one('h2.ui-search-item__title') or
                card.select_one('a.ui-search-link__title-card') or
                card.select_one('h2 a')
            )

            if not title_elem:
                return None

            titulo = title_elem.get_text(strip=True)
            url = title_elem.get('href', '')

            if not titulo or not url:
                return None

            # Precio
            price_elem = (
                card.select_one('span.andes-money-amount__fraction') or
                card.select_one('span.price-tag-fraction') or
                card.select_one('.poly-price__current .andes-money-amount')
            )

            precio, moneda = None, "CLP"
            if price_elem:
                # También buscar el símbolo de moneda
                currency_elem = card.select_one('span.andes-money-amount__currency-symbol')
                price_text = price_elem.get_text(strip=True)
                if currency_elem:
                    price_text = currency_elem.get_text(strip=True) + " " + price_text
                precio, moneda = self.clean_price(price_text)

            # Ubicación
            location_elem = (
                card.select_one('span.poly-component__location') or
                card.select_one('span.ui-search-item__location') or
                card.select_one('.ui-search-item__group__element.ui-search-item__location')
            )
            ubicacion = location_elem.get_text(strip=True) if location_elem else "Las Condes"

            # Atributos (m2, dormitorios, baños)
            metros = None
            dormitorios = None
            banos = None
            tipo_propiedad = None

            # Buscar atributos en diferentes formatos
            attrs = card.select('li.poly-attributes-list__item, li.ui-search-card-attributes__attribute')

            for attr in attrs:
                text = attr.get_text(strip=True).lower()

                if 'm²' in text or 'metros' in text:
                    metros = self.extract_number(text)
                elif 'dormitorio' in text or 'dorm' in text or 'hab' in text:
                    dormitorios = self.extract_number(text)
                elif 'baño' in text:
                    banos = self.extract_number(text)

            # También buscar en el título información del tipo
            titulo_lower = titulo.lower()
            if 'departamento' in titulo_lower or 'depto' in titulo_lower:
                tipo_propiedad = "Departamento"
            elif 'casa' in titulo_lower:
                tipo_propiedad = "Casa"
            elif 'oficina' in titulo_lower:
                tipo_propiedad = "Oficina"
            elif 'local' in titulo_lower:
                tipo_propiedad = "Local"
            elif 'estudio' in titulo_lower or 'studio' in titulo_lower:
                tipo_propiedad = "Estudio"

            # Imagen
            img_elem = (
                card.select_one('img.poly-component__picture') or
                card.select_one('img.ui-search-result-image__element') or
                card.select_one('img[data-src]')
            )
            imagen_url = None
            if img_elem:
                imagen_url = img_elem.get('data-src') or img_elem.get('src')

            return Propiedad(
                titulo=titulo,
                precio=precio,
                moneda=moneda,
                ubicacion=ubicacion,
                metros_cuadrados=metros,
                dormitorios=dormitorios,
                banos=banos,
                url=url,
                imagen_url=imagen_url,
                tipo_propiedad=tipo_propiedad,
                gastos_comunes=None,
                fecha_scraping=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            )

        except Exception as e:
            logger.warning(f"Error parseando propiedad: {e}")
            return None

    def scrape(self, max_pages: int = 3) -> List[Propiedad]:
        """
        Ejecuta el scraping de propiedades.

        Args:
            max_pages: Número máximo de páginas a procesar

        Returns:
            Lista de propiedades encontradas
        """
        self.propiedades = []
        logger.info(f"Iniciando scraping de arriendos en Las Condes ({max_pages} páginas)")

        for page in range(1, max_pages + 1):
            logger.info(f"Procesando página {page}/{max_pages}")

            url = self.build_search_url(page)
            soup = self.fetch_page(url)

            if not soup:
                logger.warning(f"No se pudo obtener la página {page}")
                continue

            # Buscar tarjetas de propiedades
            cards = soup.select('li.ui-search-layout__item')

            if not cards:
                cards = soup.select('div.ui-search-result__wrapper')

            if not cards:
                cards = soup.select('li.poly-card')

            logger.info(f"Encontradas {len(cards)} tarjetas en página {page}")

            for card in cards:
                propiedad = self.parse_property(card)
                if propiedad and propiedad.precio:
                    self.propiedades.append(propiedad)
                    logger.debug(f"Propiedad agregada: {propiedad.titulo[:50]}...")

            # Verificar si hay más páginas
            if not cards or len(cards) < 10:
                logger.info("No hay más resultados, finalizando")
                break

            if page < max_pages:
                self.random_delay()

        logger.info(f"Scraping completado: {len(self.propiedades)} propiedades encontradas")
        return self.propiedades


def export_to_excel(propiedades: List[Propiedad], filename: str = "arriendos_las_condes.xlsx"):
    """
    Exporta las propiedades a un archivo Excel con formato.

    Args:
        propiedades: Lista de propiedades a exportar
        filename: Nombre del archivo Excel
    """
    if not propiedades:
        logger.warning("No hay propiedades para exportar")
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "Arriendos Las Condes"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Headers
    headers = [
        "Título", "Precio", "Moneda", "Ubicación", "M²",
        "Dormitorios", "Baños", "Tipo", "Gastos Comunes", "URL", "Fecha Scraping"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border

    # Datos
    for row, prop in enumerate(propiedades, 2):
        ws.cell(row=row, column=1, value=prop.titulo)
        ws.cell(row=row, column=2, value=prop.precio)
        ws.cell(row=row, column=3, value=prop.moneda)
        ws.cell(row=row, column=4, value=prop.ubicacion)
        ws.cell(row=row, column=5, value=prop.metros_cuadrados)
        ws.cell(row=row, column=6, value=prop.dormitorios)
        ws.cell(row=row, column=7, value=prop.banos)
        ws.cell(row=row, column=8, value=prop.tipo_propiedad)
        ws.cell(row=row, column=9, value=prop.gastos_comunes)
        ws.cell(row=row, column=10, value=prop.url)
        ws.cell(row=row, column=11, value=prop.fecha_scraping)

        # Aplicar bordes
        for col in range(1, 12):
            ws.cell(row=row, column=col).border = border

    # Ajustar anchos de columna
    column_widths = [60, 12, 8, 25, 8, 12, 8, 15, 15, 80, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    # Formato de precio
    for row in range(2, len(propiedades) + 2):
        cell = ws.cell(row=row, column=2)
        if cell.value:
            cell.number_format = '#,##0'

    # Congelar primera fila
    ws.freeze_panes = 'A2'

    # Agregar hoja de resumen
    ws_summary = wb.create_sheet("Resumen")
    ws_summary['A1'] = "Resumen de Arriendos - Las Condes"
    ws_summary['A1'].font = Font(bold=True, size=14)

    ws_summary['A3'] = "Total Propiedades:"
    ws_summary['B3'] = len(propiedades)

    # Contar por tipo
    tipos = {}
    for p in propiedades:
        tipo = p.tipo_propiedad or "Sin especificar"
        tipos[tipo] = tipos.get(tipo, 0) + 1

    ws_summary['A5'] = "Por Tipo de Propiedad:"
    ws_summary['A5'].font = Font(bold=True)

    row = 6
    for tipo, count in sorted(tipos.items(), key=lambda x: -x[1]):
        ws_summary[f'A{row}'] = tipo
        ws_summary[f'B{row}'] = count
        row += 1

    # Precios promedio por moneda
    precios_clp = [p.precio for p in propiedades if p.moneda == "CLP" and p.precio]
    precios_uf = [p.precio for p in propiedades if p.moneda == "UF" and p.precio]

    row += 1
    ws_summary[f'A{row}'] = "Precios Promedio:"
    ws_summary[f'A{row}'].font = Font(bold=True)
    row += 1

    if precios_clp:
        ws_summary[f'A{row}'] = "CLP Promedio:"
        ws_summary[f'B{row}'] = round(sum(precios_clp) / len(precios_clp))
        ws_summary[f'B{row}'].number_format = '$#,##0'
        row += 1

    if precios_uf:
        ws_summary[f'A{row}'] = "UF Promedio:"
        ws_summary[f'B{row}'] = round(sum(precios_uf) / len(precios_uf), 2)
        row += 1

    # Guardar
    wb.save(filename)
    logger.info(f"Archivo Excel guardado: {filename}")
    return filename


def main():
    """Función principal."""
    print("=" * 60)
    print("SCRAPER DE ARRIENDOS - LAS CONDES, MERCADOLIBRE CHILE")
    print("=" * 60)

    # Crear scraper
    scraper = MercadoLibrePropiedadesScraper(
        delay_min=2,
        delay_max=4,
        max_retries=3
    )

    # Ejecutar scraping (3 páginas = ~144 propiedades máx)
    propiedades = scraper.scrape(max_pages=3)

    if propiedades:
        print(f"\nPropiedades encontradas: {len(propiedades)}")

        # Exportar a Excel
        filename = export_to_excel(propiedades, "arriendos_las_condes.xlsx")

        print(f"\nArchivo Excel generado: {filename}")

        # Mostrar resumen
        print("\n" + "=" * 60)
        print("RESUMEN")
        print("=" * 60)

        # Por tipo
        tipos = {}
        for p in propiedades:
            tipo = p.tipo_propiedad or "Sin especificar"
            tipos[tipo] = tipos.get(tipo, 0) + 1

        print("\nPor tipo de propiedad:")
        for tipo, count in sorted(tipos.items(), key=lambda x: -x[1]):
            print(f"  - {tipo}: {count}")

        # Rango de precios
        precios_clp = [p.precio for p in propiedades if p.moneda == "CLP" and p.precio]
        precios_uf = [p.precio for p in propiedades if p.moneda == "UF" and p.precio]

        if precios_clp:
            print(f"\nPrecios en CLP:")
            print(f"  - Mínimo: ${min(precios_clp):,.0f}")
            print(f"  - Máximo: ${max(precios_clp):,.0f}")
            print(f"  - Promedio: ${sum(precios_clp)/len(precios_clp):,.0f}")

        if precios_uf:
            print(f"\nPrecios en UF:")
            print(f"  - Mínimo: {min(precios_uf):.2f} UF")
            print(f"  - Máximo: {max(precios_uf):.2f} UF")
            print(f"  - Promedio: {sum(precios_uf)/len(precios_uf):.2f} UF")

    else:
        print("\nNo se encontraron propiedades.")

    return propiedades


if __name__ == "__main__":
    propiedades = main()
